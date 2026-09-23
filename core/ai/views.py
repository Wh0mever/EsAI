from django.conf import settings
from django.db import transaction
from django.http import StreamingHttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.response import Response
from rest_framework_simplejwt.authentication import JWTAuthentication
from .handlers import OpenAIHandler, stream_chat
from .models import (
    AiProvider,
    AiChatSession,
    AiChatMessage,
    AiMessageFeedback,
    AiPromptTemplate,
    AiPromptTemplateFavorite,
    GPTModels,
)
from .serializers import (
    AiChatSessionListSerializer,
    AiChatMessageSerializer,
    AiPromptTemplateSerializer,
    AiProviderSerializer,
    AiModelSerializer,
)
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction, IntegrityError

import logging

from rest_framework import status, permissions
from rest_framework.throttling import AnonRateThrottle, UserRateThrottle
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from django.utils.text import get_valid_filename
import secrets
import uuid
import os

logger = logging.getLogger(__name__)


# Максимальная длина извлечённого из документа текста (символы), чтобы не
# раздувать контекст модели и не упираться в лимит токенов OpenRouter.
DOCUMENT_TEXT_MAX_CHARS = 200_000

# Лимит размера одного загружаемого файла (байты). Переопределяется через
# settings.MAX_ATTACHMENT_SIZE. Защита от загрузки огромных файлов.
MAX_ATTACHMENT_SIZE = getattr(settings, "MAX_ATTACHMENT_SIZE", 15 * 1024 * 1024)  # 15 МБ

# Лимит количества файлов на один запрос загрузки. Переопределяется через
# settings.MAX_ATTACHMENTS_PER_REQUEST.
MAX_ATTACHMENTS_PER_REQUEST = getattr(settings, "MAX_ATTACHMENTS_PER_REQUEST", 10)

# Лимит суммарного РАСПАКОВАННОГО размера содержимого OOXML-архива (docx/pptx/
# xlsx — это zip). Защита от zip-bomb: маленький файл может распаковаться в
# гигабайты. Переопределяется через settings.MAX_OOXML_UNCOMPRESSED_SIZE.
MAX_OOXML_UNCOMPRESSED_SIZE = getattr(
    settings, "MAX_OOXML_UNCOMPRESSED_SIZE", 200 * 1024 * 1024)  # 200 МБ

# Расширения документов, из которых мы умеем извлекать текст.
_DOCUMENT_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt",
}

# OOXML-форматы (zip-контейнеры) — для них проверяем распакованный размер.
_OOXML_EXTENSIONS = {".docx", ".xlsx", ".pptx"}


def _truncate_text(text: str, limit: int = DOCUMENT_TEXT_MAX_CHARS) -> str:
    """Обрезать текст до limit символов (best-effort, не падать на None)."""
    if not text:
        return ""
    if len(text) <= limit:
        return text
    return text[:limit]


def extract_text_from_file(file, filename: str) -> str:
    """Извлечь текст из загруженного файла по его расширению.

    Поддержка:
      - .pdf            → pypdf
      - .docx           → python-docx
      - .xlsx           → openpyxl
      - .pptx           → python-pptx
      - .txt            → decode (utf-8 с фолбэком)
      - .doc/.xls/.ppt  → старые бинарные форматы НЕ поддержаны (best-effort:
                          пробуем декодировать как текст, иначе пустая строка)

    Возвращает строку, обрезанную до DOCUMENT_TEXT_MAX_CHARS символов. Никогда не
    бросает исключение — при ошибке возвращает "" (загрузка файла не должна
    падать из-за проблем извлечения текста).
    """
    name = (filename or "").lower()
    ext = os.path.splitext(name)[1]

    def _read_bytes() -> bytes:
        try:
            file.seek(0)
        except Exception:
            pass
        try:
            return file.read() or b""
        except Exception:
            return b""

    def _ooxml_uncompressed_ok(raw: bytes) -> bool:
        """Проверить РАСПАКОВАННЫЙ размер OOXML-архива до парсинга (анти zip-bomb).

        OOXML (docx/xlsx/pptx) — это zip. Маленький файл может распаковываться в
        гигабайты. Суммируем ZipInfo.file_size и отклоняем при превышении лимита.
        При ошибке чтения архива возвращаем False (безопасно — не парсим).
        """
        import io
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                total = 0
                for info in zf.infolist():
                    total += info.file_size
                    if total > MAX_OOXML_UNCOMPRESSED_SIZE:
                        return False
            return True
        except Exception:
            return False

    # Защита от zip-bomb / тяжёлых файлов: отсекаем по размеру ДО парсинга.
    # extract_text никогда не бросает наружу, поэтому при превышении просто
    # возвращаем "" (текст не извлечён, но загрузка файла не падает).
    try:
        fsize = getattr(file, "size", None)
        if fsize is not None and fsize > MAX_ATTACHMENT_SIZE:
            return ""
    except Exception:
        pass

    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(_read_bytes()))
            parts = []
            for page in reader.pages:
                try:
                    parts.append(page.extract_text() or "")
                except Exception:
                    continue
            return _truncate_text("\n".join(parts).strip())

        if ext == ".docx":
            from docx import Document
            import io
            raw = _read_bytes()
            # Анти zip-bomb: проверяем распакованный размер до парсинга.
            if not _ooxml_uncompressed_ok(raw):
                return ""
            doc = Document(io.BytesIO(raw))
            parts = []
            _total = 0  # текущая длина накопленного текста (для ранней обрезки)
            for p in doc.paragraphs:
                if p.text:
                    parts.append(p.text)
                    _total += len(p.text)
                    if _total > DOCUMENT_TEXT_MAX_CHARS:
                        return _truncate_text("\n".join(parts).strip())
            # Таблицы тоже несут текст.
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text for c in row.cells if c.text]
                    if cells:
                        line = "\t".join(cells)
                        parts.append(line)
                        _total += len(line)
                        if _total > DOCUMENT_TEXT_MAX_CHARS:
                            return _truncate_text("\n".join(parts).strip())
            return _truncate_text("\n".join(parts).strip())

        if ext == ".xlsx":
            from openpyxl import load_workbook
            import io
            raw = _read_bytes()
            # Анти zip-bomb: проверяем распакованный размер до парсинга.
            if not _ooxml_uncompressed_ok(raw):
                return ""
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts = []
            # Текущая длина — счётчиком: пересчёт sum() по всем строкам на каждой
            # новой строке делал разбор большой таблицы квадратичным.
            _total = 0
            for ws in wb.worksheets:
                title = f"# {ws.title}"
                parts.append(title)
                _total += len(title)
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        line = "\t".join(cells)
                        parts.append(line)
                        _total += len(line)
                    if _total > DOCUMENT_TEXT_MAX_CHARS:
                        break
                if _total > DOCUMENT_TEXT_MAX_CHARS:
                    break
            try:
                wb.close()
            except Exception:
                pass
            return _truncate_text("\n".join(parts).strip())

        if ext == ".pptx":
            from pptx import Presentation
            import io
            raw = _read_bytes()
            # Анти zip-bomb: проверяем распакованный размер до парсинга.
            if not _ooxml_uncompressed_ok(raw):
                return ""
            prs = Presentation(io.BytesIO(raw))
            parts = []
            _total = 0  # текущая длина накопленного текста (для ранней обрезки)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        for para in shape.text_frame.paragraphs:
                            line = "".join(run.text for run in para.runs)
                            if line:
                                parts.append(line)
                                _total += len(line)
                                if _total > DOCUMENT_TEXT_MAX_CHARS:
                                    return _truncate_text("\n".join(parts).strip())
            return _truncate_text("\n".join(parts).strip())

        if ext == ".txt":
            raw = _read_bytes()
            for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
                try:
                    return _truncate_text(raw.decode(enc).strip())
                except Exception:
                    continue
            return _truncate_text(raw.decode("utf-8", errors="ignore").strip())

        # Старые бинарные форматы .doc/.xls/.ppt — best-effort попытка декода.
        if ext in (".doc", ".xls", ".ppt"):
            raw = _read_bytes()
            try:
                decoded = raw.decode("utf-8", errors="ignore")
                # Грубая эвристика: если осмысленного текста мало — не возвращаем мусор.
                printable = sum(1 for ch in decoded if ch.isprintable() or ch in "\n\t")
                if decoded and printable / max(1, len(decoded)) > 0.6:
                    return _truncate_text(decoded.strip())
            except Exception:
                pass
            return ""
    except Exception:
        # Извлечение текста не должно ломать загрузку файла.
        return ""

    return ""


def _get_session_for_user(session_ref: str, user):
    raw = (session_ref or "").strip()
    if not raw:
        return None
    qs = AiChatSession.objects.filter(user=user)
    if raw.isdigit():
        by_id = qs.filter(id=int(raw)).first()
        if by_id:
            return by_id
    return qs.filter(uid=raw).first()


def _normalize_feedback_value(raw):
    """Привести значение фидбека к 1 / -1. None — если значение не распознано.

    Поддерживаются формы фронта: 'up'/'down', '1'/'-1', 1/-1, True/False,
    'like'/'dislike'.
    """
    if isinstance(raw, bool):
        return AiMessageFeedback.VALUE_UP if raw else AiMessageFeedback.VALUE_DOWN
    if isinstance(raw, (int, float)):
        if raw > 0:
            return AiMessageFeedback.VALUE_UP
        if raw < 0:
            return AiMessageFeedback.VALUE_DOWN
        return None
    if isinstance(raw, str):
        s = raw.strip().lower()
        if s in ("up", "like", "1", "+1", "true", "yes"):
            return AiMessageFeedback.VALUE_UP
        if s in ("down", "dislike", "-1", "false", "no"):
            return AiMessageFeedback.VALUE_DOWN
    return None


def _is_stream_request(request) -> bool:
    """True, если клиент попросил потоковый ответ (body.stream truthy)."""
    val = (getattr(request, "data", {}) or {}).get("stream")
    if isinstance(val, bool):
        return val
    if val is None:
        return False
    if isinstance(val, (int, float)):
        return val != 0
    return str(val).strip().lower() in ("1", "true", "yes", "on")


class ChatGptAPIView(APIView):
    # AllowAny: authenticated users keep the full flow (wallet/sessions/billing);
    # anonymous visitors get a limited free-quota guest flow inside OpenAIHandler.
    authentication_classes = [JWTAuthentication]
    permission_classes = [AllowAny]
    # Защита от abuse: anon/user rate-limit (ТЗ 8). Ставки в settings (anon/user).
    throttle_classes = [AnonRateThrottle, UserRateThrottle]

    def post(self, request):
        # --- Streaming path (SSE) -------------------------------------------
        # Не держим transaction.atomic вокруг всего стрима (поток долгий):
        # сохранение/списание оборачиваются в короткую транзакцию внутри
        # генератора (см. handlers.stream_chat). Пред-проверки (нет токенов →
        # 402, гость сверх лимита → 401) возвращаются обычным JSON.
        if _is_stream_request(request):
            kind, payload = stream_chat(request)
            if kind == "stream":
                resp = StreamingHttpResponse(payload, content_type="text/event-stream")
                resp["Cache-Control"] = "no-cache"
                resp["X-Accel-Buffering"] = "no"
                return resp
            # kind == "error": обычный JSON с опциональным http_status.
            http_status = 400
            if isinstance(payload, dict) and isinstance(payload.get("http_status"), int):
                http_status = payload.pop("http_status")
            return Response(payload, status=http_status)

        # --- Non-streaming path (unchanged behaviour) -----------------------
        # Сохраняем прежнюю атомарность вокруг всего обработчика, которая раньше
        # давалась декоратором @transaction.atomic на post().
        with transaction.atomic():
            response = OpenAIHandler(request)
        if isinstance(response, dict) and "error" not in response:
            return Response(response, status=200)
        # Allow the handler to request a specific HTTP status (e.g. 402 for
        # insufficient tokens). Default error status is 400.
        http_status = 400
        if isinstance(response, dict) and isinstance(response.get("http_status"), int):
            http_status = response.pop("http_status")
        return Response(response, status=http_status)


class ChatsHistory(APIView):
    
    def delete(self, request):
        sessions = AiChatSession.objects.filter(user=request.user)
        if sessions.exists():
            sessions.delete();
        return Response({}, status=200)


class ChatSessionImportAPIView(APIView):
    """Импорт гостевого диалога в аккаунт после логина (ТЗ клиента).

    Гостевой чат в БД не пишется (handlers._guest_chat — no DB writes). Когда
    гость авторизуется, фронт присылает накопленный транскрипт сюда: создаём
    сессию и сохраняем сообщения КАК ЕСТЬ — БЕЗ повторного вызова модели и БЕЗ
    списания токенов. Возвращаем uid, чтобы фронт продолжил диалог уже в сессии.
    """
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]
    throttle_classes = [UserRateThrottle]

    MAX_MESSAGES = 60
    MAX_CONTENT = 200_000

    def post(self, request):
        raw = request.data.get("messages")
        if not isinstance(raw, list) or not raw:
            return Response({"error": "no_messages"}, status=400)

        cleaned = []
        for m in raw[: self.MAX_MESSAGES * 3]:
            if not isinstance(m, dict):
                continue
            role = m.get("role")
            if role not in ("user", "assistant"):
                continue
            content = m.get("content")
            if not isinstance(content, str) or not content.strip():
                continue
            atts = m.get("attachments")
            cleaned.append({
                "role": role,
                "content": content.strip()[: self.MAX_CONTENT],
                "attachments": atts if isinstance(atts, list) else None,
            })
            if len(cleaned) >= self.MAX_MESSAGES:
                break

        if not cleaned:
            return Response({"error": "no_valid_messages"}, status=400)

        title = (request.data.get("title") or "").strip()
        if not title:
            first_user = next((c["content"] for c in cleaned if c["role"] == "user"), cleaned[0]["content"])
            title = (first_user[:60] or "Импортированный чат")

        provider = AiProvider.objects.filter(name__iexact="GPT").first()

        with transaction.atomic():
            session = AiChatSession.objects.create(
                user=request.user,
                provider=provider,
                title=title[:255],
                chat_mode="text-chat",
            )
            # Поштучно (не bulk) — чтобы created_at шёл по возрастанию и история
            # сохраняла исходный порядок при последующей загрузке сессии.
            for c in cleaned:
                AiChatMessage.objects.create(
                    session=session,
                    role=c["role"],
                    content=c["content"],
                    attachments=c["attachments"],
                    meta={"imported": True},
                )

        return Response(
            {"session": {"id": session.id, "uid": session.uid, "title": session.title, "chat_mode": session.chat_mode}},
            status=201,
        )


class ChatSessionListAPIView(APIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        # Раньше prefetch_related("messages") тянул ВСЕ сообщения всех чатов
        # (с текстами документов и вложениями), а сериализатор ещё и делал по
        # запросу на чат за последним сообщением — у активного пользователя
        # список чатов грузил десятки мегабайт. Теперь счётчик и последнее
        # сообщение считает сама база подзапросами; форма ответа прежняя.
        from django.db.models import Count, DateTimeField, IntegerField, OuterRef, Subquery
        from django.db.models.functions import Coalesce, Substr

        last = AiChatMessage.objects.filter(session=OuterRef("pk")).order_by("-created_at", "-id")
        counts = (AiChatMessage.objects.filter(session=OuterRef("pk"))
                  .order_by().values("session").annotate(n=Count("id")).values("n"))
        qs = (
            AiChatSession.objects
            .filter(user=request.user)
            .annotate(
                messages_count_value=Coalesce(Subquery(counts, output_field=IntegerField()), 0),
                last_message_role=Subquery(last.values("role")[:1]),
                # 201 символ: сериализатор режет до 200 и ставит «…», если длиннее.
                last_message_content=Subquery(last.annotate(
                    short=Substr("content", 1, 201)).values("short")[:1]),
                last_message_at=Subquery(last.values("created_at")[:1]),
            )
            # Свежая переписка — наверху: по времени последнего сообщения.
            .annotate(activity_at=Coalesce(
                "last_message_at", "updated_at", output_field=DateTimeField()))
            .order_by("-activity_at", "-id")
        )
        data = AiChatSessionListSerializer(qs, many=True).data
        return Response(data)


class ChatSessionDetailAPIView(APIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, session_ref: str):
        session = _get_session_for_user(session_ref, request.user)
        if not session:
            return Response({"error": "Session not found"}, status=404)

        messages_qs = session.messages.order_by("created_at")
        messages = AiChatMessageSerializer(messages_qs, many=True).data
        payload = {
            "id": session.id,
            "uid": session.uid,
            "title": session.title,
            "provider_id": session.provider_id,
            "created_at": session.created_at,
            "updated_at": session.updated_at,
            "is_active": session.is_active,
            "messages": messages,
        }
        return Response(payload, status=200)

    def patch(self, request, session_ref: str):
        session = _get_session_for_user(session_ref, request.user)
        if not session:
            return Response({"error": "Session not found"}, status=404)

        data = request.data or {}
        if "title" not in data:
            return Response({"detail": "Nothing to update"}, status=400)
        title = (data.get("title") or "").strip()
        if not title:
            return Response({"detail": "Title is required"}, status=400)

        session.title = title[:255]
        session.save(update_fields=["title", "updated_at"])
        return Response(
            {
                "id": session.id,
                "uid": session.uid,
                "title": session.title,
                "provider_id": session.provider_id,
                "created_at": session.created_at,
                "updated_at": session.updated_at,
                "is_active": session.is_active,
            },
            status=200,
        )

    def delete(self, request, session_ref: str):
        session = _get_session_for_user(session_ref, request.user)
        if not session:
            return Response({"error": "Session not found"}, status=404)
        session.delete()
        return Response(status=204)


class ChatMessageFeedbackAPIView(APIView):
    """POST /api/ai/chat/feedback/ — лайк/дизлайк ответа ассистента (upsert)."""
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data = request.data or {}
        message_id = data.get("message_id")
        value = _normalize_feedback_value(data.get("value"))

        if message_id is None:
            return Response({"detail": "message_id is required"}, status=400)
        if value is None:
            return Response({"detail": "Invalid value"}, status=400)

        message = AiChatMessage.objects.filter(pk=message_id).first()
        if not message:
            return Response({"error": "Message not found"}, status=404)

        # Оценивать можно только ответы в собственных сессиях пользователя.
        if message.session.user_id != request.user.id:
            return Response({"error": "Message not found"}, status=404)

        obj, _created = AiMessageFeedback.objects.update_or_create(
            message=message,
            user=request.user,
            defaults={"value": value},
        )
        return Response({"message_id": message.id, "value": obj.value}, status=200)


class AiTemplatesApiView(APIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [permissions.AllowAny]  # GET доступен всем

    def get_permissions(self):
        if self.request.method.lower() == "post":
            return [permissions.IsAuthenticated()]
        return super().get_permissions()

    def get(self, request):
        personal_data = []

        global_queryset = AiPromptTemplate.objects.filter(
            is_public=True, is_active=True
        ).select_related("provider", "created_by")
        global_data = AiPromptTemplateSerializer(global_queryset, many=True).data

        if request.user.is_authenticated:
            personal_queryset = AiPromptTemplate.objects.filter(
                is_public=False, is_active=True, created_by=request.user
            ).select_related("provider", "created_by")
            personal_data = AiPromptTemplateSerializer(personal_queryset, many=True).data

        return Response({"global": global_data, "personal": personal_data}, status=200)

    def post(self, request):
        data = request.data or {}
        name = (data.get("name") or "").strip()
        provider_id = data.get("provider")
        template = (data.get("template") or "").strip()
        purpose = (data.get("purpose") or "").strip()
        description = (data.get("description") or "").strip()
        is_active = True
        is_public = False

        # required fields
        missing = []
        if not name:
            missing.append("name")
        if not provider_id:
            missing.append("provider")
        if not template:
            missing.append("template")
        if not purpose:
            missing.append("purpose")

        if missing:
            return Response(
                {"detail": "Заполните обязательные поля", "missing": missing},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # provider check
        try:
            provider = AiProvider.objects.get(pk=provider_id)
        except AiProvider.DoesNotExist:
            return Response({"detail": "Провайдер не найден"},
                            status=status.HTTP_400_BAD_REQUEST)

        # name unique
        if AiPromptTemplate.objects.filter(name=name).exists():
            return Response({"detail": "Шаблон с таким названием уже существует"},
                            status=status.HTTP_400_BAD_REQUEST)

        try:
            with transaction.atomic():
                obj = AiPromptTemplate.objects.create(
                    name=name,
                    provider=provider,
                    template=template,
                    purpose=purpose,
                    description=description,
                    created_by=request.user,
                    is_active=is_active,
                    is_public=is_public,
                )
        except IntegrityError:
            return Response({"detail": "Не удалось создать шаблон"},
                            status=status.HTTP_400_BAD_REQUEST)

        return Response(AiPromptTemplateSerializer(obj).data,
                        status=status.HTTP_201_CREATED)
        
@method_decorator(csrf_exempt, name="dispatch")
class AiTemplatesFavoriteApiView(APIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request):
        template_id = request.data.get("templateId")
        favorite = bool(request.data.get("favorite", False))
        # TODO: persist favorite for user
        return Response({"templateId": template_id, "favorite": favorite})

    def get(self, request):
        return Response({"detail": "OK"})


class AiProvidersApiView(APIView):
    authentication_classes = []   # публичный эндпоинт
    permission_classes = []

    def get(self, request):
        qs = AiProvider.objects.all().order_by("id")
        data = AiProviderSerializer(qs, many=True).data
        return Response(data, status=status.HTTP_200_OK)


class AiModelsListAPIView(APIView):
    """Публичный список активных моделей для свитчера в чате (доступен гостям).

    JWTAuthentication опционален: токен НЕ обязателен (permission_classes пуст,
    доступ открыт всем), но если клиент прислал валидный токен — request.user
    будет реальным, и поле `locked` посчитается по его тиру (paid видит премиум
    как доступную). Без токена request.user анонимен → тир guest → премиум locked.
    """
    authentication_classes = [JWTAuthentication]
    permission_classes = []

    def get(self, request):
        qs = GPTModels.objects.filter(is_active=True).order_by("sort_order", "id")
        # context["request"] нужен сериализатору для вычисления `locked` по тиру
        # текущего пользователя (гость → премиум заблокированы).
        data = AiModelSerializer(qs, many=True, context={"request": request}).data
        return Response(data, status=status.HTTP_200_OK)


class VideoModelParamsAPIView(APIView):
    """GET /api/ai/video-params/?model=<openrouter_id> — что умеет видео-модель.

    Отдаёт список поддерживаемых длительностей (из каталога OpenRouter
    /videos/models, кэш 1 ч) и цену в кредитах за секунду, чтобы фронт мог
    показать выбор длительности со стоимостью ДО генерации.

    Зачем понадобилось: длительность была зашита в VIDEO_GEN_DURATION=5, и
    любая модель выдавала 5-секундный ролик, даже если умеет 30 (жалоба
    клиента по Seedance 2.5). Теперь длительность выбирает пользователь, а
    допустимые значения приходят отсюда — фронт не выдумывает их сам.
    """
    authentication_classes = [JWTAuthentication]
    permission_classes = []

    def get(self, request):
        model_ref = (request.query_params.get("model") or "").strip()
        if not model_ref:
            return Response({"error": "model is required"}, status=status.HTTP_400_BAD_REQUEST)

        obj = GPTModels.objects.filter(
            openrouter_id=model_ref, is_active=True, category="video"
        ).first()
        if not obj:
            return Response({"error": "video model not found"}, status=status.HTTP_404_NOT_FOUND)

        try:
            coeff = float(obj.markup_coefficient or 0)
        except Exception:
            coeff = 0.0

        durations: list[int] = []
        resolutions: list[str] = []
        aspect_ratios: list[str] = []
        sizes: list[str] = []
        frames: list[str] = []
        generate_audio = False
        try:
            from .handlers import _openrouter_video_catalog

            spec = _openrouter_video_catalog(model_ref)
            durations = [int(d) for d in (spec.get("supported_durations") or []) if str(d).isdigit()]
            resolutions = [str(x) for x in (spec.get("supported_resolutions") or []) if x]
            aspect_ratios = [str(x) for x in (spec.get("supported_aspect_ratios") or []) if x]
            sizes = [str(x) for x in (spec.get("supported_sizes") or []) if x]
            frames = [str(x) for x in (spec.get("supported_frame_images") or []) if x]
            generate_audio = bool(spec.get("generate_audio"))
        except Exception:
            logger.exception("video params lookup failed for %s", model_ref)

        default_duration = int(getattr(settings, "VIDEO_GEN_DURATION", 5) or 5)
        if durations and default_duration not in durations:
            default_duration = min(durations, key=lambda d: (abs(d - default_duration), d))

        default_resolution = str(getattr(settings, "VIDEO_GEN_RESOLUTION", "720p") or "720p")
        if resolutions and default_resolution not in resolutions:
            low = {r.lower(): r for r in resolutions}
            default_resolution = low.get(default_resolution.lower()) or resolutions[0]

        return Response({
            "model": model_ref,
            "durations": sorted(set(durations)),
            "default_duration": default_duration,
            "resolutions": resolutions,
            "default_resolution": default_resolution,
            "aspect_ratios": aspect_ratios,
            "sizes": sizes,
            # Какие кадры принимает модель: first_frame / last_frame.
            "frame_images": frames,
            "generate_audio": generate_audio,
            "credits_per_second": coeff,
        }, status=status.HTTP_200_OK)


class CartoonModelsListAPIView(APIView):
    """Модели для студии мультфильмов, сгруппированные по этапу (cartoon_stage).

    Тот же доступ и сериализация, что и ai/models/: JWTAuthentication опционален
    (permission_classes пуст, доступ открыт), а `locked` по тиру пользователя
    считается ТЕМ ЖЕ AiModelSerializer (через plans.tiers.can_use_model) — без
    токена гость видит премиум как locked, с валидным токеном — по его тиру.

    Ответ — группы {script, image, video}: активные модели с непустым
    cartoon_stage, отсортированные по sort_order. Пустые группы = [].
    """
    authentication_classes = [JWTAuthentication]
    permission_classes = []

    # Порядок и набор групп ответа = choices cartoon_stage (кроме пустого "").
    _STAGES = ("script", "image", "video")

    def get(self, request):
        qs = (
            GPTModels.objects
            .filter(is_active=True, cartoon_stage__in=self._STAGES)
            .order_by("sort_order", "id")
        )
        # context["request"] нужен сериализатору для вычисления `locked` по тиру
        # текущего пользователя — та же логика, что в ai/models/.
        grouped = {stage: [] for stage in self._STAGES}
        for model in qs:
            data = AiModelSerializer(model, context={"request": request}).data
            grouped[model.cartoon_stage].append(data)
        return Response(grouped, status=status.HTTP_200_OK)


class ChatAttachmentUploadAPIView(APIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]
    # Загрузка/парсинг документов — ресурсоёмкая операция: лимитируем (ТЗ 8).
    throttle_classes = [UserRateThrottle]

    # Расширения изображений — определяем тип по содержимому/расширению,
    # а не доверяем client-присланному content_type. SVG убран: это документ
    # со скриптами, отданный из /media/ он исполнялся бы на нашем домене.
    _IMAGE_EXTENSIONS = {
        ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic", ".heif",
    }
    # content_type → расширение для файлов без расширения (вставка скриншота
    # из буфера приходит как «image.png», но бывает и просто «blob»).
    _IMAGE_CTYPE_EXT = {
        "image/jpeg": ".jpg", "image/jpg": ".jpg", "image/png": ".png", "image/gif": ".gif",
        "image/webp": ".webp", "image/bmp": ".bmp", "image/heic": ".heic", "image/heif": ".heif",
    }

    def _validated_ext(self, f) -> str | None:
        """Расширение файла, если мы такой тип принимаем, иначе None."""
        ext = os.path.splitext((f.name or "").lower())[1]
        if ext in self._IMAGE_EXTENSIONS or ext in _DOCUMENT_EXTENSIONS:
            return ext
        if not ext:
            ctype = (getattr(f, "content_type", "") or "").lower().split(";")[0].strip()
            return self._IMAGE_CTYPE_EXT.get(ctype)
        return None

    def post(self, request):
        files = request.FILES.getlist("files")
        if not files:
            return Response({"detail": "No files provided"}, status=400)

        # Принимаем только картинки и документы, из которых умеем брать текст.
        # Раньше сохранялось что угодно под именем от клиента — .html/.svg/.js
        # из /media/ открывались браузером на нашем домене (хранимый XSS).
        for f in files:
            if self._validated_ext(f) is None:
                return Response(
                    {
                        "detail": (
                            f"Файл «{get_valid_filename(f.name or 'file')}» нельзя прикрепить. "
                            "Подходят картинки (JPG, PNG, GIF, WEBP, BMP, HEIC) и документы "
                            "(PDF, DOC, DOCX, XLS, XLSX, PPT, PPTX, TXT)."
                        ),
                        "error_code": "unsupported_file_type",
                    },
                    status=400,
                )

        # Лимит количества файлов на запрос (ТЗ 8): защита от массовой загрузки.
        if len(files) > MAX_ATTACHMENTS_PER_REQUEST:
            return Response(
                {"detail": f"Too many files (max {MAX_ATTACHMENTS_PER_REQUEST})"},
                status=400,
            )

        # Лимит размера на файл — проверяем ДО сохранения на диск.
        for f in files:
            fsize = getattr(f, "size", None)
            if fsize is not None and fsize > MAX_ATTACHMENT_SIZE:
                return Response(
                    {
                        "detail": (
                            f"File '{f.name or 'file'}' is too large "
                            f"(max {MAX_ATTACHMENT_SIZE} bytes)"
                        )
                    },
                    status=400,
                )

        uploaded = []
        for f in files:
            # Имя на диске — наше (uuid + проверенное расширение), а не клиента:
            # так в /media/ не попадёт ни «.html», ни «x.png.svg», ни путь.
            # Исходное имя остаётся только в ответе (name) — для подписи чипа.
            ext = self._validated_ext(f)
            name = get_valid_filename(f.name or "file")
            path = f"uploads/chat/{request.user.id}/{uuid.uuid4().hex}{ext}"
            saved_path = default_storage.save(path, f)
            url = default_storage.url(saved_path)
            ctype = getattr(f, "content_type", "application/octet-stream") or "application/octet-stream"

            # Тип определяем по расширению (надёжнее), а content_type клиента
            # используем лишь для файлов без расширения (см. _validated_ext).
            typ = "document" if ext in _DOCUMENT_EXTENSIONS else "image"

            item = {
                "type": typ,
                "url": url,
                "name": name,
                "size": getattr(f, "size", None),
                "content_type": ctype,
            }

            # Документы: извлекаем текст для RAG-подмешивания в контекст чата.
            if typ == "document":
                # Извлекаем из исходного UploadedFile (saved-файл уже на диске,
                # но in-memory объект ещё доступен для чтения).
                item["text"] = extract_text_from_file(f, f.name or name)

            uploaded.append(item)

        return Response({"attachments": uploaded}, status=200)


# --------------------------------------------------------------------------
# Публичные ссылки на сгенерированные ролики (esai.kg/v/<slug>)
# --------------------------------------------------------------------------

_SHARE_ALPHABET = "abcdefghijkmnopqrstuvwxyz23456789"  # без 0/1/l — их путают в переписке


def _media_name_from_url(url: str) -> str | None:
    """URL нашего медиафайла → путь в хранилище. None — если ссылка чужая.

    Та же защита, что в handlers._read_own_media: наружу принимаем только
    ссылки под MEDIA_URL, иначе пользователь смог бы «пошарить» произвольный
    адрес и наша страница стала бы витриной для чужого содержимого.
    """
    try:
        media_url = str(getattr(settings, "MEDIA_URL", "/media/") or "/media/")
        path = str(url or "")
        if path.startswith("http://") or path.startswith("https://"):
            from urllib.parse import urlparse

            path = urlparse(path).path
        if not path.startswith(media_url):
            return None
        name = path[len(media_url):].lstrip("/")
        if not name or ".." in name:
            return None
        return name
    except Exception:
        return None


class VideoShareCreateAPIView(APIView):
    """POST /api/ai/video/share/ — получить публичную ссылку на свой ролик.

    Тело: {"url": "/media/ai_videos/xxx.mp4", "prompt": "...", "model": "..."}.
    Ответ: {"slug": "...", "share_url": "https://esai.kg/v/...", ...}.

    Повторный запрос по тому же файлу отдаёт ТУ ЖЕ ссылку: человек жмёт
    «Поделиться» по нескольку раз, и плодить десяток адресов на один ролик
    (а с ними — десяток обложек) незачем.

    Обложку снимаем здесь же, при создании ссылки, а не при генерации ролика:
    делятся единицы, а генераций сотни — незачем гонять ffmpeg на каждую.
    """
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated]
    throttle_classes = [UserRateThrottle]

    def post(self, request):
        from .models import SharedVideo
        from .video_edit import extract_poster, probe_video

        raw_url = (request.data.get("url") or "").strip()
        name = _media_name_from_url(raw_url)
        if not name:
            return Response(
                {"error": "url must point to a file under MEDIA_URL"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not default_storage.exists(name):
            return Response({"error": "video not found"}, status=status.HTTP_404_NOT_FOUND)

        # Делиться можно только СВОИМ роликом. Имена файлов — uuid4 и наружу не
        # утекают, но проверять принадлежность по «его не угадаешь» нельзя:
        # достаточно один раз увидеть чужую ссылку, чтобы завести на чужое видео
        # свою страницу. Раньше искали ссылку во ВСЕХ вложениях сессий
        # пользователя — а вложения своих сообщений пишет сам клиент: вставил
        # чужую ссылку в своё сообщение — и «владеешь» роликом. Теперь верим
        # только записям сервера: журналу генераций и видео-ответам ассистента.
        from .handlers import _user_owns_video

        owns = _user_owns_video(request.user, raw_url)

        # Готовый мультфильм лежит не во вложениях чата, а в своём проекте:
        # студия делится тем же механизмом, что и ролики, поэтому право на
        # ссылку проверяем и по ней.
        kind = "video"
        if not owns:
            from cartoon.models import CartoonProject

            owns = CartoonProject.objects.filter(
                user=request.user, result_video=name).exists()
            if owns:
                kind = "cartoon"

        if not owns:
            return Response(
                {"error": "video does not belong to the current user"},
                status=status.HTTP_403_FORBIDDEN,
            )

        existing = SharedVideo.objects.filter(
            video_path=name, user=request.user, is_active=True
        ).first()
        if existing:
            return Response(self._payload(existing, request))

        prompt = (request.data.get("prompt") or "").strip()[:2000]
        model_name = (request.data.get("model") or "").strip()[:128]

        # Обложка и размеры — best-effort: без них ссылка всё равно рабочая,
        # просто превью в мессенджере будет беднее.
        poster_path, meta = "", {}
        try:
            with default_storage.open(name, "rb") as fh:
                blob = fh.read()
            meta = probe_video(blob) or {}
            poster = extract_poster(blob)
            if poster:
                poster_path = default_storage.save(
                    f"ai_videos/posters/{uuid.uuid4().hex}.jpg", ContentFile(poster)
                )
        except Exception:
            logger.exception("video share: poster/probe failed for %s", name)

        for _ in range(6):
            slug = "".join(secrets.choice(_SHARE_ALPHABET) for _ in range(10))
            try:
                share = SharedVideo.objects.create(
                    slug=slug,
                    user=request.user,
                    video_path=name,
                    poster_path=poster_path,
                    prompt=prompt,
                    model_name=model_name,
                    kind=kind,
                    width=meta.get("width"),
                    height=meta.get("height"),
                    duration=meta.get("duration"),
                )
                break
            except IntegrityError:
                continue  # астрономически редкое совпадение кода — берём новый
        else:
            return Response(
                {"error": "could not allocate share link"},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        return Response(self._payload(share, request), status=status.HTTP_201_CREATED)

    @staticmethod
    def _payload(share, request) -> dict:
        # Домен берём из самого запроса, а не из настроек: сайт отвечает и на
        # esai.kg, и на www.esai.kg — ссылка должна остаться на том домене, где
        # человек нажал «Поделиться».
        return {
            "slug": share.slug,
            "kind": share.kind,
            "share_url": request.build_absolute_uri(f"/v/{share.slug}"),
            "video_url": request.build_absolute_uri(default_storage.url(share.video_path)),
        }


class VideoSharePublicAPIView(APIView):
    """GET /api/ai/video/share/<slug>/ — данные ролика для публичной страницы.

    Без авторизации: страницу открывают по ссылке из мессенджера, где никакой
    сессии нет. Отдаём только сам ролик и запрос к модели — ни автора, ни его
    почты здесь быть не должно.
    """
    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [AnonRateThrottle]

    def get(self, request, slug):
        from .models import SharedVideo

        share = SharedVideo.objects.filter(slug=slug, is_active=True).first()
        if not share:
            return Response({"error": "not found"}, status=status.HTTP_404_NOT_FOUND)

        # F() вместо share.views + 1: страницу могут открыть одновременно, а
        # чтение-запись в питоне потеряет часть просмотров.
        from django.db.models import F

        SharedVideo.objects.filter(pk=share.pk).update(views=F("views") + 1)

        def abs_media(path: str) -> str:
            return request.build_absolute_uri(default_storage.url(path)) if path else ""

        return Response({
            "slug": share.slug,
            "kind": share.kind,
            "video_url": abs_media(share.video_path),
            "poster_url": abs_media(share.poster_path),
            "prompt": share.prompt,
            "model": share.model_name,
            "width": share.width,
            "height": share.height,
            "duration": share.duration,
            "created_at": share.created_at.isoformat(),
        })
